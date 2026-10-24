from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _write_text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _parse_datetime(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return str(value)


def export_financial_workbook(orders, output_path, returns=(), store_name="EarthGlow"):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Ringkasan"
    transactions = workbook.create_sheet("Transaksi")
    return_sheet = workbook.create_sheet("Retur")

    header_fill = PatternFill("solid", fgColor="286344")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=16, bold=True, color="286344")
    currency_format = '"Rp" #,##0;[Red]-"Rp" #,##0'

    summary["A1"] = f"Laporan Keuangan {store_name}"
    summary["A1"].font = title_font
    summary_headers = [
        "Saluran",
        "Transaksi lunas",
        "Omzet",
        "HPP tercatat",
        "Laba kotor",
        "Uang kembali",
        "Nilai kerugian retur",
        "Laba berkurang karena retur",
        "Laba setelah retur",
        "Transaksi tanpa HPP",
        "Retur laba belum diketahui",
    ]
    for column, header in enumerate(summary_headers, start=1):
        cell = summary.cell(row=3, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row_number, title in (
        (4, "Transaksi langsung"),
        (5, "Pesanan online"),
        (6, "TOTAL"),
    ):
        summary.cell(row=row_number, column=1, value=title)

    summary["A8"] = (
        "Laba kotor = omzet - HPP. Retur memperhitungkan nilai pengembalian "
        "atau biaya barang pengganti; laba setelah retur tidak termasuk biaya "
        "operasional. Transaksi lama tanpa HPP tidak dianggap bermodal nol."
    )
    summary["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A8:K9")
    summary.row_dimensions[8].height = 34
    summary.row_dimensions[9].height = 26
    for column, width in {
        "A": 26, "B": 18, "C": 20, "D": 20, "E": 20,
        "F": 20, "G": 22, "H": 28, "I": 22, "J": 22, "K": 26,
    }.items():
        summary.column_dimensions[column].width = width

    headers = [
        "ID",
        "Saluran",
        "Tanggal",
        "Pelanggan",
        "Produk",
        "Jumlah",
        "Metode bayar",
        "Status",
        "Omzet (Rp)",
        "HPP (Rp)",
        "Laba (Rp)",
    ]
    for column, header in enumerate(headers, start=1):
        cell = transactions.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_number, order in enumerate(orders, start=2):
        order_id, customer, _, product, quantity, total, payment, status, created_at, cost_total, channel = order
        values = [
            order_id,
            "direct" if channel == "direct" else "online",
            _parse_datetime(created_at),
            customer or "Pelanggan umum",
            product or "-",
            quantity or 0,
            payment or "Cash",
            status or "Belum Lunas",
            total or 0,
            cost_total,
        ]
        for column, value in enumerate(values, start=1):
            cell = transactions.cell(row=row_number, column=column)
            if column in (3, 9, 10):
                cell.value = value
            else:
                _write_text(cell, value)
            if column == 3 and isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            if column in (9, 10):
                cell.number_format = currency_format

        profit_cell = transactions.cell(row=row_number, column=11)
        profit_cell.value = (
            f'=IF(H{row_number}<>"Lunas",0,'
            f'IF(J{row_number}="","",I{row_number}-J{row_number}))'
        )
        profit_cell.number_format = currency_format

    return_headers = [
        "Saluran", "Tanggal", "Nama pemesan", "Barang retur", "Jumlah",
        "Keluhan", "Nilai barang retur", "HPP barang retur", "Cara ganti rugi",
        "Uang kembali", "Barang pengganti", "HPP pengganti",
        "Nilai kerugian", "Laba berkurang",
    ]
    for column, header in enumerate(return_headers, start=1):
        cell = return_sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row_number, record in enumerate(returns, start=2):
        (
            _, _, channel, customer, product, quantity, complaint,
            returned_price, returned_cost, resolution, refund_amount, replacement,
            replacement_cost, _, _, created_at
        ) = record
        values = [
            channel,
            _parse_datetime(created_at),
            customer or "Pelanggan umum",
            product,
            quantity,
            complaint,
            returned_price,
            returned_cost,
            resolution,
            refund_amount,
            replacement or "",
            replacement_cost,
        ]
        for column, value in enumerate(values, start=1):
            cell = return_sheet.cell(row=row_number, column=column)
            if column in (2, 7, 8, 10, 12):
                cell.value = value
            else:
                _write_text(cell, value)
            if column == 2 and isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            if column in (7, 8, 10, 12):
                cell.number_format = currency_format
        return_sheet.cell(
            row=row_number,
            column=13,
            value=f'=IF(I{row_number}="Uang kembali",J{row_number},L{row_number})',
        ).number_format = currency_format
        return_sheet.cell(
            row=row_number,
            column=14,
            value=(
                f'=IF(AND(I{row_number}="Uang kembali",H{row_number}=""),"",'
                f'IF(I{row_number}="Uang kembali",G{row_number}-H{row_number},'
                f'MAX(L{row_number}-IF(H{row_number}="",0,H{row_number}),0)))'
            ),
        ).number_format = currency_format

    order_last = max(transactions.max_row, 2)
    return_last = max(return_sheet.max_row, 2)
    channel_range = f"'Transaksi'!$B$2:$B${order_last}"
    status_range = f"'Transaksi'!$H$2:$H${order_last}"
    revenue_range = f"'Transaksi'!$I$2:$I${order_last}"
    cost_range = f"'Transaksi'!$J$2:$J${order_last}"
    profit_range = f"'Transaksi'!$K$2:$K${order_last}"
    return_channel_range = f"'Retur'!$A$2:$A${return_last}"
    refund_range = f"'Retur'!$J$2:$J${return_last}"
    loss_range = f"'Retur'!$M$2:$M${return_last}"
    return_profit_range = f"'Retur'!$N$2:$N${return_last}"
    for row_number, channel in ((4, "direct"), (5, "online")):
        criterion = f'"{channel}"'
        summary.cell(
            row=row_number,
            column=2,
            value=f'=COUNTIFS({channel_range},{criterion},{status_range},"Lunas")',
        )
        summary.cell(
            row=row_number,
            column=3,
            value=f'=SUMIFS({revenue_range},{channel_range},{criterion},{status_range},"Lunas")',
        )
        summary.cell(
            row=row_number,
            column=4,
            value=f'=SUMIFS({cost_range},{channel_range},{criterion},{status_range},"Lunas")',
        )
        summary.cell(
            row=row_number,
            column=5,
            value=f'=SUMIFS({profit_range},{channel_range},{criterion})',
        )
        summary.cell(
            row=row_number,
            column=6,
            value=f'=SUMIFS({refund_range},{return_channel_range},{criterion})',
        )
        summary.cell(
            row=row_number,
            column=7,
            value=f'=SUMIFS({loss_range},{return_channel_range},{criterion})',
        )
        summary.cell(
            row=row_number,
            column=8,
            value=f'=SUMIFS({return_profit_range},{return_channel_range},{criterion})',
        )
        summary.cell(row=row_number, column=9, value=f"=E{row_number}-H{row_number}")
        summary.cell(
            row=row_number,
            column=10,
            value=f'=COUNTIFS({channel_range},{criterion},{status_range},"Lunas",{cost_range},"")',
        )
        summary.cell(
            row=row_number,
            column=11,
            value=f'=COUNTIFS({return_channel_range},{criterion},{return_profit_range},"")',
        )
    for column in range(2, 12):
        letter = chr(ord("A") + column - 1)
        summary.cell(row=6, column=column, value=f"=SUM({letter}4:{letter}5)")
    for row_number in (4, 5, 6):
        summary.cell(row=row_number, column=2).number_format = "0"
        summary.cell(row=row_number, column=10).number_format = "0"
        summary.cell(row=row_number, column=11).number_format = "0"
        for column in range(3, 10):
            summary.cell(row=row_number, column=column).number_format = currency_format

    for column, width in {
        "A": 10,
        "B": 18,
        "C": 20,
        "D": 32,
        "E": 12,
        "F": 24,
        "G": 20,
        "H": 20,
        "I": 20,
        "J": 20,
        "K": 20,
    }.items():
        transactions.column_dimensions[column].width = width
    transactions.freeze_panes = "A2"
    transactions.auto_filter.ref = f"A1:K{transactions.max_row}"
    for column, width in enumerate((16, 20, 24, 30, 12, 40, 20, 20, 20, 20, 30, 20, 20, 20), start=1):
        return_sheet.column_dimensions[chr(ord("A") + column - 1)].width = width
    return_sheet.freeze_panes = "A2"
    return_sheet.auto_filter.ref = f"A1:N{return_sheet.max_row}"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
